"""RP02 Cargo Handling — backdated historical discharge upload and integration.

Prior to PORTMAN go-live (01-May-2026), vessel and MBC cargo discharge records
are not present in live operation tables (ldud_header / mbc_header / lueu_lines).
This module handles parsing, validation, template generation, storage, and retrieval
of backdated Cargo Handling data, seamlessly merged with live rows in the UI and Excel.
"""
import io
import re
import csv as _csv
from datetime import datetime, date
from database import get_db, get_cursor
from .model import _blank, _norm_header, read_matrix

TABLE = 'rp02_cargo_handling_backdated'

FIELDS = [
    ('vessel_name',         'Vessel Name'),
    ('vessel_type',         'MV/MBC'),
    ('material_po',         'Material PO'),
    ('cargo_type',          'Type'),
    ('cargo_name',          'Cargo'),
    ('bl_qty_mt',           'B/L Qty (MT)'),
    ('actual_discharge',    'Actual Discharge (MT)'),
    ('load_port',           'Load Port'),
    ('discharge_commenced', 'Discharge Commenced'),
    ('discharge_completed', 'Discharge Completed'),
    ('consignee',           'Consignee/Customer'),
    ('flag',                'Flag'),
    ('invoice_number',      'Invoice No.'),
    ('status',              'Status'),
]

COLUMNS = [f for f, _ in FIELDS]
TEMPLATE_HEADERS = ['Sr'] + [h for _, h in FIELDS]

TEMPLATE_EXAMPLE_ROWS = [
    ['1', 'MV OBE Lotus', 'MV', '4100550083', 'IBRM', 'Orissa Fines',
     '19,500.00', '19,450.00', 'Orissa', '27-03-2026 20:30', '03-04-2026 05:30',
     'JSW Steel Ltd', 'Panama', 'DPPL/25-26/101', 'Completed'],
    ['2', 'MBC Ganga', 'MBC', '4100550099', 'COAL', 'Coking Coal',
     '5,000.00', '4,980.00', 'Paradip', '05-04-2026 10:00', '08-04-2026 14:00',
     'Tata Steel Ltd', 'India', 'DPPL/25-26/102', 'Completed'],
]

ALIASES = {
    'sr':                   'sr',
    'srno':                 'sr',
    'sno':                  'sr',
    'slno':                 'sr',
    'serial':               'sr',
    'vesselname':           'vessel_name',
    'vessel':               'vessel_name',
    'ship':                 'vessel_name',
    'mvmbc':                'vessel_type',
    'vesseltype':           'vessel_type',
    'type':                 'cargo_type',
    'cargotype':            'cargo_type',
    'po':                   'material_po',
    'materialpo':           'material_po',
    'ponumber':             'material_po',
    'cargo':                'cargo_name',
    'commodity':            'cargo_name',
    'cargoname':            'cargo_name',
    'blqty':                'bl_qty_mt',
    'blqtymt':              'bl_qty_mt',
    'blquantity':           'bl_qty_mt',
    'blquantitymt':         'bl_qty_mt',
    'actualdischarge':      'actual_discharge',
    'actualdischargemt':    'actual_discharge',
    'dischargedqty':        'actual_discharge',
    'dischargeqty':         'actual_discharge',
    'loadport':             'load_port',
    'port':                 'load_port',
    'dischargecommenced':   'discharge_commenced',
    'commenced':            'discharge_commenced',
    'commenceddate':        'discharge_commenced',
    'dischargecompleted':   'discharge_completed',
    'completed':            'discharge_completed',
    'completeddate':        'discharge_completed',
    'consignee':            'consignee',
    'customer':             'consignee',
    'party':                'consignee',
    'client':               'consignee',
    'consigneecustomer':    'consignee',
    'flag':                 'flag',
    'vesselflag':           'flag',
    'invoiceno':            'invoice_number',
    'invoicenumber':        'invoice_number',
    'billno':               'invoice_number',
    'status':               'status',
}

_NUM_COLS = {'bl_qty_mt', 'actual_discharge'}
_DT_COLS = {'discharge_commenced', 'discharge_completed'}

_DT_FORMATS = [
    '%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y-%m-%d',
    '%d-%m-%Y %H:%M:%S', '%d-%m-%Y %H:%M', '%d-%m-%Y',
    '%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y',
    '%d.%m.%Y %H:%M:%S', '%d.%m.%Y %H:%M', '%d.%m.%Y',
    '%m/%d/%Y %I:%M:%S %p', '%m/%d/%Y %I:%M %p', '%m/%d/%Y',
    '%d-%b-%Y %H:%M:%S', '%d-%b-%Y %H:%M', '%d-%b-%Y',
    '%d-%b-%y %H:%M:%S', '%d-%b-%y %H:%M', '%d-%b-%y',
]

def _clean_num(v):
    if _blank(v): return None
    s = str(v).replace(',', '').replace(' ', '').strip()
    try:
        f = float(s)
        return ('%d' % int(f)) if f.is_integer() else ('%.3f' % f).rstrip('0').rstrip('.')
    except ValueError:
        return str(v).strip()

def parse_datetime(v):
    if _blank(v): return None
    if isinstance(v, datetime): return v
    if isinstance(v, date): return datetime.combine(v, datetime.min.time())
    s = str(v).strip()
    if s.lower() in ('inprogress', 'in progress', 'in-progress'): return 'InProgress'
    for fmt in _DT_FORMATS:
        try: return datetime.strptime(s, fmt)
        except ValueError: pass
    return None

def _format_clean_dt(v):
    if _blank(v): return ''
    dt = parse_datetime(v)
    if dt == 'InProgress': return 'InProgress'
    if isinstance(dt, datetime):
        if dt.hour == 0 and dt.minute == 0 and dt.second == 0:
            return dt.strftime('%Y-%m-%d')
        return dt.strftime('%Y-%m-%d %H:%M:%S')
    return str(v).strip()

def _text(v):
    if _blank(v): return None
    if isinstance(v, datetime): return v.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(v, date): return v.strftime('%Y-%m-%d')
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()

def _header_key(value):
    """
    Normalize an Excel/CSV header into a simple comparison key.
    Examples:
        Vessel Name -> vesselname
        B/L Qty (MT) -> blqtymt
        Consignee/Customer -> consigneecustomer
    """
    if value is None:
        return ''

    s = str(value).strip().lower()
    s = s.replace('&', 'and')
    s = re.sub(r'[^a-z0-9]+', '', s)
    return s


def _index(headers):
    """
    Map Excel headers to internal field names.
    Also handles slightly corrupted/missing characters in Excel headers.
    """
    from difflib import SequenceMatcher

    idx = {}

    header_aliases = dict(ALIASES)
    header_aliases.update({
        'vesselname': 'vessel_name',
        'vesselame': 'vessel_name',
        'nameofvessel': 'vessel_name',
        'shipname': 'vessel_name',

        'mvmbc': 'vessel_type',
        'mvormbc': 'vessel_type',
        'vesseltype': 'vessel_type',

        'materialpo': 'material_po',
        'maeialpo': 'material_po',
        'materialponumber': 'material_po',
        'ponumber': 'material_po',
        'po': 'material_po',

        'type': 'cargo_type',
        'ype': 'cargo_type',
        'cargotype': 'cargo_type',

        'cargo': 'cargo_name',
        'cago': 'cargo_name',
        'cargoname': 'cargo_name',
        'commodity': 'cargo_name',

        'blqty': 'bl_qty_mt',
        'blqtymt': 'bl_qty_mt',
        'blqym': 'bl_qty_mt',
        'blquantity': 'bl_qty_mt',
        'blquantitymt': 'bl_qty_mt',

        'actualdischarge': 'actual_discharge',
        'actualdischargemt': 'actual_discharge',
        'acualdischagem': 'actual_discharge',
        'acualdischargemt': 'actual_discharge',
        'actualdischarged': 'actual_discharge',
        'dischargedqty': 'actual_discharge',
        'dischargeqty': 'actual_discharge',

        'loadport': 'load_port',
        'loadpo': 'load_port',
        'portofloading': 'load_port',

        'dischargecommenced': 'discharge_commenced',
        'dischargecommence': 'discharge_commenced',
        'dischagecommenced': 'discharge_commenced',
        'dischagecommece': 'discharge_commenced',
        'commenced': 'discharge_commenced',
        'commenceddate': 'discharge_commenced',

        'dischargecompleted': 'discharge_completed',
        'dischargecomplete': 'discharge_completed',
        'dischagecompleted': 'discharge_completed',
        'dischagecompleed': 'discharge_completed',
        'completed': 'discharge_completed',
        'completeddate': 'discharge_completed',

        'consignee': 'consignee',
        'customer': 'consignee',
        'consigneecustomer': 'consignee',
        'cosigeecusome': 'consignee',
        'party': 'consignee',
        'client': 'consignee',

        'flag': 'flag',
        'vesselflag': 'flag',

        'invoiceno': 'invoice_number',
        'ivoiceo': 'invoice_number',
        'invoicenumber': 'invoice_number',
        'billno': 'invoice_number',

        'status': 'status',
    })

    canonical = {
        'vesselname': 'vessel_name',
        'vesseltype': 'vessel_type',
        'mvmbc': 'vessel_type',
        'materialpo': 'material_po',
        'type': 'cargo_type',
        'cargotype': 'cargo_type',
        'cargo': 'cargo_name',
        'cargoname': 'cargo_name',
        'blqty': 'bl_qty_mt',
        'blqtymt': 'bl_qty_mt',
        'actualdischarge': 'actual_discharge',
        'actualdischargemt': 'actual_discharge',
        'loadport': 'load_port',
        'dischargecommenced': 'discharge_commenced',
        'dischargecompleted': 'discharge_completed',
        'consignee': 'consignee',
        'consigneecustomer': 'consignee',
        'flag': 'flag',
        'invoiceno': 'invoice_number',
        'invoicenumber': 'invoice_number',
        'status': 'status',
    }

    for i, h in enumerate(headers):
        key = _header_key(h)

        if not key:
            continue

        canon = header_aliases.get(key)

        if not canon:
            try:
                canon = ALIASES.get(_norm_header(h))
            except Exception:
                canon = None

        # Fuzzy fallback handles strings such as vesselame -> vesselname.
        if not canon:
            best_key = None
            best_score = 0.0

            for candidate in canonical:
                score = SequenceMatcher(None, key, candidate).ratio()

                if key.endswith('mt') and candidate.endswith('mt'):
                    score = max(
                        score,
                        SequenceMatcher(
                            None, key[:-2], candidate[:-2]
                        ).ratio()
                    )

                if score > best_score:
                    best_score = score
                    best_key = candidate

            if best_key and best_score >= 0.68:
                canon = canonical[best_key]

        if canon and canon not in idx:
            idx[canon] = i

    return idx


def parse_rows(headers, rows):
    idx = _index(headers)
    out, errors = [], []
    for n, raw in enumerate(rows, start=2):
        if not raw or all(_blank(c) for c in raw): continue
        def get(canon):
            col = idx.get(canon)
            if col is None or col >= len(raw): return None
            return raw[col]
        vessel_name = _text(get('vessel_name'))
        if not vessel_name:
            errors.append({'row': n, 'message': 'Vessel Name is required'})
            continue
        r = {}
        for col in COLUMNS:
            val = get(col)
            if col in _NUM_COLS: r[col] = _clean_num(val)
            elif col in _DT_COLS: r[col] = _format_clean_dt(val)
            else: r[col] = _text(val)
        out.append(r)
    return out, errors

def parse_upload(file_storage):
    """
    Read uploaded CSV/XLSX and locate the Cargo Handling header row.
    """
    matrix, errors = read_matrix(file_storage)

    if errors:
        return [], errors

    if not matrix:
        return [], [{
            'row': 1,
            'message': 'The uploaded file is empty.'
        }]

    for i, row in enumerate(matrix):
        if not row:
            continue

        raw_headers = [
            str(c).strip() if c is not None else ''
            for c in row
        ]

        idx = _index(raw_headers)

        if (
            'vessel_name' in idx
            and (
                len(idx) >= 3
                or 'actual_discharge' in idx
                or 'bl_qty_mt' in idx
            )
        ):
            parsed_rows, row_errors = parse_rows(
                raw_headers,
                matrix[i + 1:]
            )

            header_excel_row = i + 1

            for err in row_errors:
                if isinstance(err.get('row'), int):
                    err['row'] = header_excel_row + (err['row'] - 1)

            return parsed_rows, row_errors

    detected = []

    for i, row in enumerate(matrix[:20]):
        keys = [
            _header_key(c)
            for c in row
            if c is not None and str(c).strip()
        ]

        if keys:
            detected.append(
                f"Row {i + 1}: " + ", ".join(keys[:20])
            )

    return [], [{
        'row': 0,
        'message': (
            'Could not locate a valid Cargo Handling header row. '
            'Expected Vessel Name and other Cargo Handling columns. '
            'Detected rows: ' + ' | '.join(detected[:5])
        )
    }]


def build_template_csv():
    buf = io.StringIO()
    w = _csv.writer(buf)
    w.writerow(TEMPLATE_HEADERS)
    for r in TEMPLATE_EXAMPLE_ROWS: w.writerow(r)
    return buf.getvalue()

def build_template_excel():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Cargo_Handling_Template'
    thin = Side(style='thin', color='CBD5E0')
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill('solid', fgColor='E8E7B5')
    hdr_font = Font(name='Arial', size=11, bold=True, color='000000')
    align_ctr = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center')
    ws.append(TEMPLATE_HEADERS)
    for col_idx in range(1, len(TEMPLATE_HEADERS) + 1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = hdr_fill
        c.font = hdr_font
        c.border = bdr
        c.alignment = align_ctr
    for row_data in TEMPLATE_EXAMPLE_ROWS: ws.append(row_data)
    for r_idx in range(2, len(TEMPLATE_EXAMPLE_ROWS) + 2):
        for c_idx in range(1, len(TEMPLATE_HEADERS) + 1):
            c = ws.cell(row=r_idx, column=c_idx)
            c.border = bdr
            c.font = Font(name='Arial', size=10)
            c.alignment = align_ctr if c_idx in (1, 3, 14) else align_left
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()

def replace_all(rows, uploaded_by):
    conn = get_db()
    cur = get_cursor(conn)
    try:
        cur.execute('DELETE FROM %s' % TABLE)
        cols = COLUMNS + ['uploaded_by']
        sql = 'INSERT INTO %s (%s) VALUES (%s)' % (TABLE, ', '.join(cols), ', '.join(['%s'] * len(cols)))
        for r in rows:
            cur.execute(sql, [r.get(c) for c in COLUMNS] + [str(uploaded_by) if uploaded_by else None])
        conn.commit()
        return len(rows)
    except Exception:
        conn.rollback()
        raise
    finally: conn.close()

def append_rows(rows, uploaded_by):
    conn = get_db()
    cur = get_cursor(conn)
    try:
        cols = COLUMNS + ['uploaded_by']
        sql = 'INSERT INTO %s (%s) VALUES (%s)' % (TABLE, ', '.join(cols), ', '.join(['%s'] * len(cols)))
        for r in rows:
            cur.execute(sql, [r.get(c) for c in COLUMNS] + [str(uploaded_by) if uploaded_by else None])
        conn.commit()
        return len(rows)
    except Exception:
        conn.rollback()
        raise
    finally: conn.close()

def get_status():
    conn = get_db()
    cur = get_cursor(conn)
    try:
        cur.execute('SELECT COUNT(*) AS c, MAX(uploaded_at) AS at FROM %s' % TABLE)
        r = cur.fetchone() or {'c': 0, 'at': None}
        cur.execute('SELECT COUNT(DISTINCT vessel_name) AS v FROM %s' % TABLE)
        v_res = cur.fetchone() or {'v': 0}
        at = r['at']
        return {
            'count': r['c'],
            'uploaded_at': at.strftime('%Y-%m-%d %H:%M') if at else None,
            'vessels': v_res['v']
        }
    finally: conn.close()

def get_rows(page=1, size=50, filters=None):
    conn = get_db()
    cur = get_cursor(conn)
    try:
        where, params = [], []
        if filters:
            for f in filters:
                field, val = f.get('field'), f.get('value')
                if field and val and field in COLUMNS:
                    where.append(f"{field} ILIKE %s")
                    params.append(f"%{val}%")
        where_clause = ('WHERE ' + ' AND '.join(where)) if where else ''
        cur.execute(f"SELECT COUNT(*) AS total FROM {TABLE} {where_clause}", params)
        total = cur.fetchone()['total']
        offset = (page - 1) * size
        cur.execute(f"SELECT id, {', '.join(COLUMNS)}, uploaded_at FROM {TABLE} {where_clause} ORDER BY id DESC LIMIT %s OFFSET %s", params + [size, offset])
        rows = cur.fetchall()
        out = []
        for r in rows:
            d = dict(r)
            if d.get('uploaded_at'): d['uploaded_at'] = d['uploaded_at'].strftime('%Y-%m-%d %H:%M')
            out.append(d)
        return out, total
    finally: conn.close()

def update_row(row_id, fields_dict):
    allowed = {k: v for k, v in fields_dict.items() if k in COLUMNS}
    if not allowed: return {'error': 'No editable fields provided'}
    conn = get_db()
    cur = get_cursor(conn)
    try:
        set_parts = [f"{col} = %s" for col in allowed.keys()]
        sql = f"UPDATE {TABLE} SET {', '.join(set_parts)} WHERE id = %s"
        cur.execute(sql, list(allowed.values()) + [row_id])
        conn.commit()
        return {'success': True}
    except Exception as e:
        conn.rollback()
        return {'error': str(e)}
    finally: conn.close()

def delete_row(row_id):
    conn = get_db()
    cur = get_cursor(conn)
    try:
        cur.execute(f"DELETE FROM {TABLE} WHERE id = %s", [row_id])
        conn.commit()
    finally: conn.close()

def _parse_filter_date(d_str):
    if not d_str: return None
    try: return datetime.strptime(d_str.strip()[:10], '%Y-%m-%d').date()
    except (ValueError, TypeError): return None

def get_backdated_cargo_rows(from_date=None, to_date=None):
    conn = get_db()
    cur = get_cursor(conn)
    try:
        cur.execute(f"SELECT * FROM {TABLE} ORDER BY id ASC")
        raw_rows = cur.fetchall()
        f_date = _parse_filter_date(from_date)
        t_date = _parse_filter_date(to_date)
        out = []
        for r in raw_rows:
            commenced_dt = parse_datetime(r.get('discharge_commenced'))
            completed_dt = parse_datetime(r.get('discharge_completed'))
            comm_date = commenced_dt.date() if isinstance(commenced_dt, datetime) else None
            comp_date = completed_dt.date() if isinstance(completed_dt, datetime) else None
            if f_date or t_date:
                active_date = comm_date or comp_date
                if not active_date: continue
                if f_date and t_date:
                    start_d = comm_date or comp_date
                    end_d = comp_date or comm_date
                    if not (start_d <= t_date and end_d >= f_date): continue
                elif f_date and (comm_date or comp_date) < f_date: continue
                elif t_date and (comm_date or comp_date) > t_date: continue
            def _fmt_ui_dt(dt_obj, raw_val):
                if isinstance(dt_obj, datetime): return dt_obj.strftime('%d/%m/%Y %H:%M')
                return str(raw_val or '').strip()
            def _fmt_num(val):
                if _blank(val): return 0.0
                try: return float(str(val).replace(',', '').strip())
                except (ValueError, TypeError): return 0.0
            out.append({
                'id': f"bd_{r['id']}",
                'customer_detail_id': 0,
                'vessel_name': r.get('vessel_name') or '',
                'vessel_type': r.get('vessel_type') or '',
                'material_po': r.get('material_po') or '',
                'cargo_type': r.get('cargo_type') or '',
                'cargo_name': r.get('cargo_name') or '',
                'bl_qty_mt': _fmt_num(r.get('bl_qty_mt')),
                'actual_discharge': _fmt_num(r.get('actual_discharge')),
                'load_port': r.get('load_port') or '',
                'consignee': r.get('consignee') or '',
                'discharge_commenced': _fmt_ui_dt(commenced_dt, r.get('discharge_commenced')),
                'discharge_completed': _fmt_ui_dt(completed_dt, r.get('discharge_completed')),
                'flag': r.get('flag') or '',
                'invoice_number': r.get('invoice_number') or '',
                'status': r.get('status') or 'Completed',
                'source': 'Backdated'
            })
        return out
    finally: conn.close()
